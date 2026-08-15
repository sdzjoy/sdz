import csv
import hashlib
import io
import re
import zipfile
from pathlib import Path

from django.conf import settings
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.validators import URLValidator
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from standards.models import Standard

from .models import ImportBatch, ImportRow, Resource, ResourceMirror, ResourceVersion

MAX_IMPORT_ROWS = 500
MAX_XLSX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_XLSX_ENTRIES = 1000

HEADER_ALIASES = {
    "title": "title",
    "名称": "title",
    "资源名称": "title",
    "slug": "slug",
    "网址标识": "slug",
    "summary": "summary",
    "简介": "summary",
    "access_level": "access_level",
    "访问等级": "access_level",
    "copyright_status": "copyright_status",
    "版权状态": "copyright_status",
    "copyright_note": "copyright_note",
    "版权说明": "copyright_note",
    "author": "author",
    "作者": "author",
    "source_name": "source_name",
    "来源名称": "source_name",
    "source_url": "source_url",
    "来源网址": "source_url",
    "file_format": "file_format",
    "文件格式": "file_format",
    "size_bytes": "size_bytes",
    "大小_字节": "size_bytes",
    "checksum_algorithm": "checksum_algorithm",
    "校验算法": "checksum_algorithm",
    "checksum_value": "checksum_value",
    "校验值": "checksum_value",
    "version": "version",
    "版本": "version",
    "standard_codes": "standard_codes",
    "相关标准编号": "standard_codes",
    "mirror_provider": "mirror_provider",
    "网盘": "mirror_provider",
    "share_url": "share_url",
    "分享链接": "share_url",
    "extraction_code": "extraction_code",
    "提取码": "extraction_code",
}

ACCESS_ALIASES = {
    "L0": Resource.AccessLevel.PUBLIC,
    "公开": Resource.AccessLevel.PUBLIC,
    "公开区": Resource.AccessLevel.PUBLIC,
    "L1": Resource.AccessLevel.MEMBER,
    "普通会员": Resource.AccessLevel.MEMBER,
    "普通会员区": Resource.AccessLevel.MEMBER,
    "L2": Resource.AccessLevel.ADVANCED,
    "高级会员": Resource.AccessLevel.ADVANCED,
    "高级会员区": Resource.AccessLevel.ADVANCED,
    "L3": Resource.AccessLevel.OWNER,
    "站长": Resource.AccessLevel.OWNER,
    "私人区": Resource.AccessLevel.OWNER,
}

COPYRIGHT_ALIASES = {
    value: value for value in Resource.CopyrightStatus.values
} | {label: value for value, label in Resource.CopyrightStatus.choices}

PROVIDER_ALIASES = {
    value: value for value in ResourceMirror.Provider.values
} | {label: value for value, label in ResourceMirror.Provider.choices}

FORBIDDEN_HEADER_PARTS = {
    "password",
    "密码",
    "cookie",
    "token",
    "oauth",
    "账号",
    "账户",
    "username",
    "用户名",
    "客户端密钥",
}

RESOURCE_IMPORT_FIELDS = {
    "title",
    "slug",
    "summary",
    "access_level",
    "copyright_status",
    "copyright_note",
    "author",
    "source_name",
    "source_url",
    "file_format",
    "size_bytes",
    "checksum_algorithm",
    "checksum_value",
}


def _require_staff(actor):
    if actor is None or not actor.is_active or not actor.is_staff:
        raise PermissionDenied("只有管理员可以预览或确认资源导入。")


def _clean_header(value):
    return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")


def _validate_headers(headers):
    cleaned = [_clean_header(header) for header in headers]
    for header in cleaned:
        if any(part in header for part in FORBIDDEN_HEADER_PARTS):
            raise ValidationError("导入文件包含账号、密码、Cookie 或令牌类禁用字段。")
    unknown = [header for header in cleaned if header and header not in HEADER_ALIASES]
    if unknown:
        raise ValidationError("导入文件包含未知列：" + "、".join(unknown))
    normalized = [HEADER_ALIASES.get(header, "") for header in cleaned]
    duplicates = {header for header in normalized if header and normalized.count(header) > 1}
    if duplicates:
        raise ValidationError("导入文件包含重复含义的列：" + "、".join(sorted(duplicates)))
    return normalized


def _rows_from_csv(data):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValidationError("CSV 必须使用 UTF-8 编码。") from exc
    reader = csv.reader(io.StringIO(text))
    try:
        raw_headers = next(reader)
    except StopIteration as exc:
        raise ValidationError("导入文件为空。") from exc
    headers = _validate_headers(raw_headers)
    rows = []
    for row_number, values in enumerate(reader, start=2):
        if not any(str(value).strip() for value in values):
            continue
        if len(rows) >= MAX_IMPORT_ROWS:
            raise ValidationError(f"单次最多导入 {MAX_IMPORT_ROWS} 行。")
        values = list(values) + [""] * max(0, len(headers) - len(values))
        rows.append((row_number, dict(zip(headers, values, strict=False))))
    return rows


def _validate_xlsx_archive(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise ValidationError("Excel 内部文件数量超过安全上限。")
            expanded_size = sum(entry.file_size for entry in entries)
            if expanded_size > MAX_XLSX_EXPANDED_BYTES:
                raise ValidationError("Excel 解压后大小超过安全上限。")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ValidationError("不支持加密的 Excel 文件。")
    except zipfile.BadZipFile as exc:
        raise ValidationError("Excel 文件结构无效。") from exc


def _rows_from_xlsx(data):
    _validate_xlsx_archive(data)
    try:
        workbook = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise ValidationError("无法读取 Excel 文件。") from exc
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(iterator)
        except StopIteration as exc:
            raise ValidationError("导入文件为空。") from exc
        headers = _validate_headers(raw_headers)
        rows = []
        for row_number, values in enumerate(iterator, start=2):
            if not any(str(value or "").strip() for value in values):
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                raise ValidationError(f"单次最多导入 {MAX_IMPORT_ROWS} 行。")
            values = list(values) + [""] * max(0, len(headers) - len(values))
            rows.append((row_number, dict(zip(headers, values, strict=False))))
        return rows
    finally:
        workbook.close()


def _string(value, limit=None):
    if value is None:
        return ""
    text = str(value).strip()
    return text[:limit] if limit else text


def _positive_integer(value, errors):
    if value in (None, ""):
        return None
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        errors.append("size_bytes 必须是非负整数。")
        return None
    if parsed < 0:
        errors.append("size_bytes 必须是非负整数。")
        return None
    return parsed


def _standard_ids(value, errors):
    codes = [
        item.strip().upper()
        for item in re.split(r"[|,，;；]", _string(value))
        if item.strip()
    ]
    if not codes:
        return []
    found = dict(Standard.objects.filter(code__in=codes).values_list("code", "pk"))
    missing = [code for code in codes if code not in found]
    if missing:
        errors.append("相关标准尚未录入：" + "、".join(missing))
    return [found[code] for code in codes if code in found]


def _normalize_row(raw, seen_slugs):
    errors = []
    normalized = {
        "title": _string(raw.get("title"), 300),
        "slug": _string(raw.get("slug"), 160),
        "summary": _string(raw.get("summary"), 5000),
        "access_level": ACCESS_ALIASES.get(_string(raw.get("access_level")), ""),
        "copyright_status": COPYRIGHT_ALIASES.get(
            _string(raw.get("copyright_status")),
            "",
        ),
        "copyright_note": _string(raw.get("copyright_note"), 2000),
        "author": _string(raw.get("author"), 300),
        "source_name": _string(raw.get("source_name"), 300),
        "source_url": _string(raw.get("source_url"), 1000),
        "file_format": _string(raw.get("file_format"), 120),
        "size_bytes": _positive_integer(raw.get("size_bytes"), errors),
        "checksum_algorithm": _string(raw.get("checksum_algorithm"), 20),
        "checksum_value": _string(raw.get("checksum_value"), 256),
        "version": _string(raw.get("version"), 120),
        "standard_ids": _standard_ids(raw.get("standard_codes"), errors),
        "mirror_provider": PROVIDER_ALIASES.get(_string(raw.get("mirror_provider")), ""),
        "share_url": _string(raw.get("share_url"), 1500),
        "extraction_code": _string(raw.get("extraction_code"), 80),
    }
    if not normalized["title"]:
        errors.append("资源名称不能为空。")
    if not normalized["slug"]:
        errors.append("网址标识不能为空。")
    elif normalized["slug"] in seen_slugs:
        errors.append("同一文件中的网址标识重复。")
    else:
        seen_slugs.add(normalized["slug"])
    if normalized["slug"] and Resource.objects.filter(slug=normalized["slug"]).exists():
        errors.append("网址标识已存在。")
    if not normalized["summary"]:
        errors.append("简介不能为空。")
    if not normalized["access_level"]:
        errors.append("访问等级必须是 L0、L1、L2 或 L3。")
    if not normalized["copyright_status"]:
        errors.append("版权状态无效或未填写。")
    if normalized["mirror_provider"] and not normalized["share_url"]:
        errors.append("填写网盘类型时必须填写分享链接。")
    if normalized["share_url"] and not normalized["mirror_provider"]:
        errors.append("填写分享链接时必须填写网盘类型。")
    for field_name in ("source_url", "share_url"):
        if normalized[field_name]:
            try:
                URLValidator()(normalized[field_name])
            except ValidationError:
                errors.append(f"{field_name} 不是有效网址。")

    resource_values = {name: normalized[name] for name in RESOURCE_IMPORT_FIELDS}
    resource = Resource(
        **resource_values,
        status=Resource.Status.PUBLISHED,
        published_at=timezone.now(),
    )
    try:
        resource.full_clean(validate_unique=False)
    except ValidationError as exc:
        errors.extend(exc.messages)
    return normalized, list(dict.fromkeys(errors))


def preview_import(*, uploaded_file, actor):
    _require_staff(actor)
    filename = Path(uploaded_file.name).name[:255]
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise ValidationError("只支持 CSV 或 XLSX 文件。")
    data = uploaded_file.read(settings.FILE_UPLOAD_MAX_MEMORY_SIZE + 1)
    if len(data) > settings.FILE_UPLOAD_MAX_MEMORY_SIZE:
        raise ValidationError("导入文件超过 5 MiB 上限。")
    batch = ImportBatch.objects.create(
        filename=filename,
        file_type=extension.removeprefix("."),
        file_hash=hashlib.sha256(data).hexdigest(),
        state=ImportBatch.State.FAILED,
        created_by=actor,
    )
    try:
        parsed_rows = _rows_from_csv(data) if extension == ".csv" else _rows_from_xlsx(data)
        if not parsed_rows:
            raise ValidationError("导入文件没有数据行。")
        seen_slugs = set()
        invalid_count = 0
        for row_number, raw in parsed_rows:
            raw = {key: _string(value) for key, value in raw.items() if key}
            normalized, errors = _normalize_row(raw, seen_slugs)
            state = ImportRow.State.INVALID if errors else ImportRow.State.VALID
            invalid_count += int(bool(errors))
            ImportRow.objects.create(
                batch=batch,
                row_number=row_number,
                raw_data=raw,
                normalized_data=normalized,
                errors=errors,
                state=state,
            )
        batch.row_count = len(parsed_rows)
        batch.error_count = invalid_count
        batch.state = ImportBatch.State.INVALID if invalid_count else ImportBatch.State.READY
        batch.save(update_fields=("row_count", "error_count", "state"))
    except ValidationError as exc:
        batch.state = ImportBatch.State.FAILED
        batch.error_message = "；".join(exc.messages)[:1000]
        batch.save(update_fields=("state", "error_message"))
    return batch


@transaction.atomic
def confirm_import(*, batch, actor):
    _require_staff(actor)
    locked = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    rows = list(locked.rows.select_for_update().order_by("row_number"))
    if locked.state != ImportBatch.State.READY:
        raise ValidationError("只有全部行校验通过的预览批次才能确认发布。")
    if not rows or any(row.state != ImportRow.State.VALID for row in rows):
        raise ValidationError("导入行状态已变化，请重新预览后再确认。")

    now = timezone.now()
    for row in rows:
        normalized = dict(row.normalized_data)
        standard_ids = normalized.pop("standard_ids", [])
        version = normalized.pop("version", "")
        provider = normalized.pop("mirror_provider", "")
        share_url = normalized.pop("share_url", "")
        extraction_code = normalized.pop("extraction_code", "")
        resource_values = {
            name: normalized.get(name) for name in RESOURCE_IMPORT_FIELDS
        }
        resource = Resource(
            **resource_values,
            status=Resource.Status.PUBLISHED,
            published_at=now,
        )
        resource.full_clean()
        resource.save()
        if standard_ids:
            resource.related_standards.add(*standard_ids)
        if version:
            ResourceVersion.objects.create(
                resource=resource,
                version=version,
                file_format=resource.file_format,
                size_bytes=resource.size_bytes,
                checksum_algorithm=resource.checksum_algorithm,
                checksum_value=resource.checksum_value,
            )
        if share_url:
            ResourceMirror.objects.create(
                resource=resource,
                provider=provider,
                share_url=share_url,
                extraction_code=extraction_code,
                last_verified_at=now,
                last_verified_by=actor,
            )
        row.created_resource = resource
        row.state = ImportRow.State.IMPORTED
        row.normalized_data = {
            **normalized,
            "standard_ids": standard_ids,
            "version": version,
            "mirror_provider": provider,
            "share_url": "[已转入资源入口]" if share_url else "",
            "extraction_code": "[已转入资源入口]" if extraction_code else "",
        }
        raw_data = dict(row.raw_data)
        if raw_data.get("share_url"):
            raw_data["share_url"] = "[已转入资源入口]"
        if raw_data.get("extraction_code"):
            raw_data["extraction_code"] = "[已转入资源入口]"
        row.raw_data = raw_data
        row.save(
            update_fields=("created_resource", "state", "normalized_data", "raw_data")
        )

    locked.state = ImportBatch.State.CONFIRMED
    locked.confirmed_by = actor
    locked.confirmed_at = now
    locked.save(update_fields=("state", "confirmed_by", "confirmed_at"))
    batch.state = locked.state
    batch.confirmed_by = actor
    batch.confirmed_at = now
    return locked
